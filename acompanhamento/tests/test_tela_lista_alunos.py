"""Tela de alunos por situação no marco e exportação para Excel."""

from datetime import date
from io import BytesIO
from unittest.mock import patch
from zipfile import ZipFile

from django.test import TestCase, override_settings
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from acompanhamento.lista_alunos import (
    IRREGULARES,
    MARCO_INTERMEDIARIO,
    QUALQUER_MARCO,
    cartao_alunos_irregulares,
)
from acompanhamento.models import MarcoCalendario, PlanilhaBaixada
from acompanhamento.tests.test_lista_alunos import criar_aluno
from acompanhamento.tests.test_linha_de_base import criar_perfis


@override_settings(SECURE_SSL_REDIRECT=False, ACOMPANHAMENTO_ATIVO=True)
class TelaListaAlunosTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.perfis = criar_perfis()
        cls.endereco = reverse("coordenacao_alunos")
        MarcoCalendario.objects.create(
            semestre="2025.1",
            tipo=MarcoCalendario.TipoMarco.PROJETO_DISSERTACAO,
            data_limite=date(2025, 6, 1),
        )
        criar_aluno(
            "Atrasado Tela Marco",
            prazo_qualificacao="2025.1",
            data_limite_qualificacao=date(2025, 6, 1),
        )

    def setUp(self):
        self.client.force_login(self.perfis["servidor"])

    def test_lista_filtrada_exibe_situacao_e_botao_excel(self):
        resposta = self.client.get(
            self.endereco,
            {
                "nome": "Atrasado Tela Marco",
                "marco": MARCO_INTERMEDIARIO,
                "situacao_marco": IRREGULARES,
            },
        )

        self.assertEqual(resposta.status_code, 200)
        self.assertContains(resposta, "Atrasado Tela Marco")
        self.assertContains(resposta, "Irregular")
        self.assertNotContains(resposta, "com chance")
        self.assertNotContains(resposta, "sem chance")
        self.assertContains(resposta, "Baixar em Excel")
        filtros = {
            filtro["titulo"]: filtro["valor"]
            for filtro in resposta.context["filtros_ativos"]
        }
        self.assertEqual(filtros["Situação no marco"], "Irregular")

    def test_total_do_cartao_e_o_da_lista_aberta_pelo_endereco(self):
        cartao = cartao_alunos_irregulares()
        resposta = self.client.get(cartao["endereco"])

        self.assertGreaterEqual(cartao["total"], 1)
        self.assertEqual(resposta.context["pagina"].paginator.count, cartao["total"])

    def test_opcoes_de_situacao_no_marco(self):
        resposta = self.client.get(self.endereco, {"marco": MARCO_INTERMEDIARIO})

        self.assertEqual(
            resposta.context["situacoes_marco"],
            (
                ("irregulares", "Irregular"),
                ("regular", "Regular"),
                ("nao_calculado", "Não foi possível calcular"),
            ),
        )

    def test_opcao_qualquer_marco_e_coluna_indicam_os_marcos_irregulares(self):
        criar_aluno(
            "Irregular nos Dois Marcos",
            data_ingresso=date(2020, 1, 1),
            data_limite_qualificacao=date(2025, 5, 1),
            prazo_qualificacao="2025.1",
        )
        resposta = self.client.get(
            self.endereco,
            {
                "nome": "Irregular nos Dois Marcos",
                "marco": QUALQUER_MARCO,
                "situacao_marco": IRREGULARES,
            },
        )

        self.assertEqual(resposta.status_code, 200)
        self.assertEqual(resposta.context["marcos"][-1], (QUALQUER_MARCO, "Qualquer marco"))
        self.assertEqual(resposta.context["pagina"].paginator.count, 1)
        self.assertContains(resposta, "Projeto de dissertação ou qualificação e Defesa")
        self.assertContains(resposta, "Situação no marco")

    def test_colunas_do_marco_vem_logo_depois_do_aluno(self):
        resposta = self.client.get(
            self.endereco, {"marco": MARCO_INTERMEDIARIO, "situacao_marco": IRREGULARES}
        )
        conteudo = resposta.content.decode()

        self.assertLess(conteudo.index(">Situação no marco</th>"), conteudo.index(">Matrícula</th>"))
        self.assertContains(resposta, '<span class="badge badge-no">Irregular</span>', html=True)

    def test_sem_marco_a_lista_nao_mostra_as_colunas_do_marco(self):
        resposta = self.client.get(self.endereco)

        self.assertNotContains(resposta, ">Situação no marco</th>")
        self.assertContains(resposta, ">Matrícula</th>")

    def test_qualquer_marco_recusa_situacao_ou_semestre_incompativeis(self):
        for parametros, mensagem in (
            (
                {"marco": QUALQUER_MARCO},
                '“Qualquer marco” só aceita a situação “Irregular”.',
            ),
            (
                {"marco": QUALQUER_MARCO, "situacao_marco": "regular"},
                '“Qualquer marco” só aceita a situação “Irregular”.',
            ),
            (
                {
                    "marco": QUALQUER_MARCO,
                    "situacao_marco": IRREGULARES,
                    "semestre_marco": "2025.1",
                },
                'Não é possível filtrar “Qualquer marco” por semestre.',
            ),
        ):
            with self.subTest(parametros=parametros):
                resposta = self.client.get(self.endereco, parametros)
                self.assertEqual(resposta.status_code, 200)
                self.assertEqual(
                    [str(item) for item in resposta.context["messages"]],
                    [mensagem],
                )

    @override_settings(ACOMPANHAMENTO_ATIVO=False)
    def test_parametros_de_marco_sao_ignorados_com_modulo_desligado(self):
        sem_parametros = self.client.get(self.endereco)
        com_marco = self.client.get(
            self.endereco,
            {
                "marco": MARCO_INTERMEDIARIO,
                "situacao_marco": "invalida",
                "semestre_marco": "2025-1",
                "exportar": "xlsx",
            },
        )

        self.assertEqual(sem_parametros.status_code, 200)
        self.assertEqual(com_marco.status_code, 200)
        self.assertEqual(
            [aluno.pk for aluno in sem_parametros.context["alunos"]],
            [aluno.pk for aluno in com_marco.context["alunos"]],
        )
        for campo in ("marco", "situacao_marco", "semestre_marco"):
            self.assertNotContains(com_marco, f'name="{campo}"')
        self.assertNotContains(com_marco, "Baixar em Excel")
        for coluna in (
            "Situação no marco",
            "Data limite",
            "Reapresentação até",
            "Semestre do marco",
            "Detalhe",
        ):
            self.assertNotContains(com_marco, coluna)
        self.assertNotEqual(
            com_marco["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

    def test_excel_contem_apenas_alunos_correspondentes_aos_filtros(self):
        inicio_download = timezone.now()
        resposta = self.client.get(
            self.endereco,
            {
                "nome": "Atrasado Tela Marco",
                "marco": MARCO_INTERMEDIARIO,
                "situacao_marco": IRREGULARES,
                "exportar": "xlsx",
            },
        )

        self.assertEqual(resposta.status_code, 200)
        self.assertEqual(
            resposta["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        workbook = load_workbook(BytesIO(resposta.content), read_only=True)
        linhas = list(workbook.active.values)

        self.assertEqual(linhas[6][0], "ALUNO")
        self.assertEqual(linhas[7][0], "Atrasado Tela Marco")
        self.assertEqual(linhas[7][1], "Irregular")
        self.assertEqual(len(linhas), 8)
        self.assertIn(
            f'alunos_por_marco_{MARCO_INTERMEDIARIO}_{timezone.localdate().isoformat()}.xlsx',
            resposta["Content-Disposition"],
        )
        registro = PlanilhaBaixada.objects.get()
        self.assertEqual(registro.baixada_por, self.perfis["servidor"])
        self.assertGreaterEqual(registro.baixada_em, inicio_download)
        self.assertLessEqual(registro.baixada_em, timezone.now())
        self.assertEqual(
            registro.filtros,
            {
                "nome": "Atrasado Tela Marco",
                "marco": MARCO_INTERMEDIARIO,
                "situacao_marco": IRREGULARES,
            },
        )

    def test_excel_qualquer_marco_tem_os_marcos_irregulares_na_linha(self):
        criar_aluno(
            "Excel Irregular nos Dois",
            data_ingresso=date(2020, 1, 1),
            data_limite_qualificacao=date(2025, 5, 1),
            prazo_qualificacao="2025.1",
        )
        resposta = self.client.get(
            self.endereco,
            {
                "nome": "Excel Irregular nos Dois",
                "marco": QUALQUER_MARCO,
                "situacao_marco": IRREGULARES,
                "exportar": "xlsx",
            },
        )

        self.assertEqual(resposta.status_code, 200)
        workbook = load_workbook(BytesIO(resposta.content), read_only=True)
        linhas = list(workbook.active.values)
        self.assertEqual(len(linhas), 8)
        self.assertEqual(linhas[7][1], "Projeto de dissertação ou qualificação e Defesa")
        self.assertEqual(linhas[7][2], "Irregular")

    def _baixar_planilha(self):
        return self.client.get(
            self.endereco,
            {
                "nome": "Atrasado Tela Marco",
                "marco": MARCO_INTERMEDIARIO,
                "situacao_marco": IRREGULARES,
                "exportar": "xlsx",
            },
        )

    def test_excel_tem_o_topo_e_as_cores_do_acadflow(self):
        planilha = load_workbook(BytesIO(self._baixar_planilha().content)).active

        self.assertEqual(
            planilha["A2"].value, "PPGEC · Pós-Graduação em Engenharia de Computação"
        )
        self.assertEqual(planilha["A3"].value, "Alunos por marco acadêmico")
        self.assertIn("Nome: Atrasado Tela Marco", planilha["A4"].value)
        self.assertIn("Situação no marco: Irregular", planilha["A4"].value)
        self.assertIn("1 aluno ", planilha["A5"].value)
        self.assertEqual(planilha["A1"].fill.start_color.rgb[-6:], "002F69")
        self.assertEqual(planilha["A7"].font.color.rgb[-6:], "00489A")
        self.assertEqual(planilha["B8"].font.color.rgb[-6:], "B42336")
        self.assertEqual(planilha["B8"].fill.start_color.rgb[-6:], "FFF0F1")
        self.assertEqual(planilha.freeze_panes, "A8")
        self.assertEqual(planilha.auto_filter.ref, "A7:L8")
        self.assertFalse(planilha.sheet_view.showGridLines)

    def test_excel_traz_a_logo_e_a_data_como_data(self):
        arquivo = ZipFile(BytesIO(self._baixar_planilha().content))

        self.assertIn("xl/media/image1.png", arquivo.namelist())
        planilha = load_workbook(BytesIO(self._baixar_planilha().content)).active
        self.assertEqual(planilha["C8"].value.date(), date(2025, 6, 1))
        self.assertEqual(planilha["C8"].number_format, "DD/MM/YYYY")

    def test_excel_sem_a_logo_continua_sendo_baixado(self):
        with patch("acompanhamento.planilha_alunos.finders.find", return_value="nao-existe.png"):
            with self.assertLogs("acompanhamento.planilha_alunos", level="ERROR"):
                resposta = self._baixar_planilha()

        self.assertEqual(resposta.status_code, 200)
        self.assertEqual(
            load_workbook(BytesIO(resposta.content)).active["A8"].value,
            "Atrasado Tela Marco",
        )

    def test_excel_grava_nome_iniciado_com_igual_como_texto(self):
        criar_aluno(
            "=1+1",
            prazo_qualificacao="2025.1",
            data_limite_qualificacao=date(2025, 6, 1),
        )

        resposta = self.client.get(
            self.endereco,
            {"nome": "=1+1", "marco": MARCO_INTERMEDIARIO, "exportar": "xlsx"},
        )

        self.assertEqual(resposta.status_code, 200)
        workbook = load_workbook(BytesIO(resposta.content))
        celula_nome = workbook.active["A8"]
        self.assertEqual(celula_nome.value, "=1+1")
        self.assertEqual(celula_nome.data_type, "s")

    def test_filtro_de_semestre_invalido_exibe_erro(self):
        resposta = self.client.get(
            self.endereco,
            {
                "marco": MARCO_INTERMEDIARIO,
                "semestre_marco": "2025-1",
            },
        )

        self.assertEqual(resposta.status_code, 200)
        self.assertContains(resposta, "Semestre do marco inválido.")
