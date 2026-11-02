"""Planilha da lista de alunos por marco, com a identidade visual do Acadflow.

As cores vêm de static/css/app.css. As colunas seguem a ordem e os nomes da tela.
"""

import logging
import math
import struct
from io import BytesIO
from pathlib import Path

from django.contrib.staticfiles import finders
from django.http import HttpResponse
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.drawing.image import Image
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.utils.units import pixels_to_EMU
from openpyxl.worksheet.properties import PageSetupProperties

from .lista_alunos import QUALQUER_MARCO

logger = logging.getLogger(__name__)

# Tokens de static/css/app.css
COR_MARCA = "002F69"  # --color-primary-strong
COR_PRIMARIA = "00489A"  # --color-primary
COR_PRIMARIA_SUAVE = "EAF3FF"  # --color-primary-soft
COR_TEXTO = "18233A"  # --color-text
COR_APOIO = "637083"  # --color-muted
COR_BORDA = "E3E9F2"  # --color-border
COR_SUBTITULO_MARCA = "C9DCF5"  # borda de .badge-info

# Situação: (fundo, texto), de .badge-no, .badge-ok e .badge-aviso
CORES_SITUACAO = {
    "Irregular": ("FFF0F1", "B42336"),
    "Regular": ("EAF8F0", "14643D"),
}
CORES_SITUACAO_SEM_CALCULO = ("FFF6E9", "92510A")

FONTE = "Segoe UI"
LOGO = "img/acadflow-wordmark-branco.png"
ALTURA_LOGO_PX = 30
ALTURA_FAIXA_PT = 40
PROGRAMA = "PPGEC · Pós-Graduação em Engenharia de Computação"
SEM_VALOR = "—"

LINHA_CABECALHO = 7
COLUNA_ALUNO = 1
TITULO_MARCO = "Marco"
TITULO_SITUACAO = "Situação no marco"
# (título, largura, centralizada, de apoio)
COLUNAS = (
    ("Aluno", 32, False, False),
    (TITULO_MARCO, 26, False, False),
    (TITULO_SITUACAO, 27, True, False),
    ("Data limite", 25, True, False),
    ("Reapresentação até", 25, True, False),
    ("Semestre do marco", 24, True, True),
    ("Detalhe", 58, False, True),
    ("Matrícula", 14, False, True),
    ("E-mail", 44, False, True),
    ("Polo", 16, False, True),
    ("Nível", 13, False, True),
    ("Ingresso", 11, True, True),
    ("Orientador", 28, False, True),
)


class _Logo(Image):
    """Imagem PNG sem exigir o Pillow, que o projeto não instala.

    A classe do openpyxl só usa o Pillow para descobrir o tamanho e o formato.
    O restante da escrita usa apenas os atributos definidos aqui.
    """

    def __init__(self, dados):
        largura, altura = struct.unpack(">II", dados[16:24])
        self.ref = BytesIO(dados)
        self._dados = dados
        self.format = "png"
        self.height = ALTURA_LOGO_PX
        self.width = round(largura * ALTURA_LOGO_PX / altura)
        tamanho = XDRPositiveSize2D(pixels_to_EMU(self.width), pixels_to_EMU(self.height))
        margem_superior = round((ALTURA_FAIXA_PT * 96 / 72 - ALTURA_LOGO_PX) / 2)
        self.anchor = OneCellAnchor(
            _from=AnchorMarker(
                col=0,
                colOff=pixels_to_EMU(12),
                row=0,
                rowOff=pixels_to_EMU(margem_superior),
            ),
            ext=tamanho,
        )

    def _data(self):
        return self._dados


def _colocar_logo(planilha):
    """Uma logo que não abre não pode impedir o download da planilha."""
    caminho = finders.find(LOGO)
    try:
        planilha.add_image(_Logo(Path(caminho).read_bytes()))
    except (OSError, TypeError, struct.error):
        logger.exception("Não foi possível colocar a logo na planilha de alunos.")


def _fonte(**opcoes):
    return Font(name=FONTE, size=opcoes.pop("size", 10), **opcoes)


def _fundo(cor):
    return PatternFill("solid", start_color=cor, end_color=cor)


def _colunas_do_marco(marco):
    """Como na tela, a coluna Marco só aparece em "Qualquer marco"."""
    if marco == QUALQUER_MARCO:
        return COLUNAS
    return tuple(coluna for coluna in COLUNAS if coluna[0] != TITULO_MARCO)


def _altura_da_linha(valores, colunas):
    """Altura com respiro, como as linhas da tabela da tela."""
    linhas = 1
    for valor, (_, largura, *_resto) in zip(valores, colunas):
        if isinstance(valor, str):
            necessarias = sum(
                max(1, math.ceil(len(parte) / (largura * 1.08))) for parte in valor.split("\n")
            )
            linhas = max(linhas, necessarias)
    return max(24, linhas * 14 + 10)


def _texto_por_marco(linha, campo, formatar=lambda valor: valor):
    return "\n".join(
        f"{nome}: {formatar(getattr(resultado, campo)) if getattr(resultado, campo) else SEM_VALOR}"
        for nome, resultado in linha.resultados_marco
    )


def _data_br(data):
    return data.strftime("%d/%m/%Y")


def _valores_do_aluno(linha, marco):
    resultado = linha.resultado
    if marco == QUALQUER_MARCO:
        marco_da_linha = linha.marco_display
        data_limite = _texto_por_marco(linha, "data_limite", _data_br)
        reapresentacao = _texto_por_marco(linha, "prazo_reapresentacao", _data_br)
        detalhe = _texto_por_marco(linha, "motivo")
    else:
        marco_da_linha = None
        data_limite = resultado.data_limite
        reapresentacao = resultado.prazo_reapresentacao
        detalhe = resultado.motivo
    valores = [
        linha.aluno.nome,
        *([marco_da_linha] if marco == QUALQUER_MARCO else []),
        linha.situacao_display,
        data_limite,
        reapresentacao,
        linha.semestre_marco,
        detalhe,
        linha.aluno.matricula,
        linha.aluno.email,
        linha.aluno.polo_atuacao.nome if linha.aluno.polo_atuacao else "",
        linha.trajetoria.get_nivel_curso_display(),
        linha.trajetoria.ingresso,
        linha.trajetoria.orientador.nome if linha.trajetoria.orientador else "",
    ]
    return [valor if valor else SEM_VALOR for valor in valores]


def _escrever_topo(planilha, total, filtros, colunas):
    ultima_coluna = len(colunas)
    for coluna in range(1, ultima_coluna + 1):
        planilha.cell(row=1, column=coluna).fill = _fundo(COR_MARCA)
        planilha.cell(row=2, column=coluna).fill = _fundo(COR_MARCA)
    planilha.row_dimensions[1].height = ALTURA_FAIXA_PT
    _colocar_logo(planilha)

    agora = timezone.localtime()
    filtros_texto = "  |  ".join(f"{filtro['titulo']}: {filtro['valor']}" for filtro in filtros)
    resumo = f"{total} aluno{'' if total == 1 else 's'}  |  Gerada em {agora:%d/%m/%Y às %H:%M}"
    topo = (
        (2, PROGRAMA, _fonte(size=9, color=COR_SUBTITULO_MARCA), 20),
        (3, "Alunos por marco acadêmico", _fonte(size=16, bold=True, color=COR_TEXTO), 34),
        (4, f"Filtros: {filtros_texto}" if filtros_texto else "Sem filtros", _fonte(color=COR_APOIO), 18),
        (5, resumo, _fonte(color=COR_APOIO), 18),
        (6, None, _fonte(), 8),
    )
    for linha, texto, fonte, altura in topo:
        planilha.merge_cells(start_row=linha, start_column=1, end_row=linha, end_column=ultima_coluna)
        celula = planilha.cell(row=linha, column=1, value=texto)
        celula.font = fonte
        celula.alignment = Alignment(vertical="center", indent=1)
        planilha.row_dimensions[linha].height = altura


def _escrever_cabecalho(planilha, colunas):
    linha_inferior = Side(style="medium", color=COR_PRIMARIA)
    for coluna, (titulo, largura, centralizada, _apoio) in enumerate(colunas, start=1):
        planilha.column_dimensions[get_column_letter(coluna)].width = largura
        celula = planilha.cell(row=LINHA_CABECALHO, column=coluna, value=titulo.upper())
        celula.font = _fonte(size=9, bold=True, color=COR_PRIMARIA)
        celula.fill = _fundo(COR_PRIMARIA_SUAVE)
        celula.border = Border(bottom=linha_inferior)
        celula.alignment = Alignment(
            horizontal="center" if centralizada else "left",
            vertical="center",
            wrap_text=True,
            indent=0 if centralizada else 1,
        )
    planilha.row_dimensions[LINHA_CABECALHO].height = 34


def _escrever_aluno(planilha, numero_da_linha, valores, colunas):
    linha_inferior = Border(bottom=Side(style="thin", color=COR_BORDA))
    for coluna, (valor, (_titulo, _largura, centralizada, apoio)) in enumerate(
        zip(valores, colunas), start=1
    ):
        celula = planilha.cell(row=numero_da_linha, column=coluna, value=valor)
        # Nome que comece com "=" é texto, nunca fórmula.
        if isinstance(valor, str) and valor.startswith("="):
            celula.data_type = "s"
        if not isinstance(valor, str):
            celula.number_format = "DD/MM/YYYY"
        celula.font = _fonte(color=COR_APOIO if apoio else COR_TEXTO)
        celula.border = linha_inferior
        celula.alignment = Alignment(
            horizontal="center" if centralizada else "left",
            vertical="center",
            wrap_text=True,
            indent=0 if centralizada else 1,
        )

    aluno = planilha.cell(row=numero_da_linha, column=COLUNA_ALUNO)
    aluno.font = _fonte(bold=True, color=COR_PRIMARIA)

    coluna_situacao = [coluna[0] for coluna in colunas].index(TITULO_SITUACAO) + 1
    situacao = planilha.cell(row=numero_da_linha, column=coluna_situacao)
    fundo, texto = CORES_SITUACAO.get(situacao.value, CORES_SITUACAO_SEM_CALCULO)
    situacao.fill = _fundo(fundo)
    situacao.font = _fonte(bold=True, color=texto)

    planilha.row_dimensions[numero_da_linha].height = _altura_da_linha(valores, colunas)


def _configurar_exibicao(planilha, colunas):
    ultima_coluna = get_column_letter(len(colunas))
    # Só as linhas: fixar uma coluna cortaria o título, que ocupa a largura toda.
    planilha.freeze_panes = planilha.cell(row=LINHA_CABECALHO + 1, column=COLUNA_ALUNO)
    planilha.auto_filter.ref = f"A{LINHA_CABECALHO}:{ultima_coluna}{planilha.max_row}"
    planilha.sheet_view.showGridLines = False
    planilha.sheet_properties.tabColor = COR_PRIMARIA
    planilha.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    planilha.page_setup.orientation = "landscape"
    planilha.page_setup.paperSize = planilha.PAPERSIZE_A4
    planilha.page_setup.fitToWidth = 1
    planilha.page_setup.fitToHeight = 0
    planilha.print_title_rows = f"{LINHA_CABECALHO}:{LINHA_CABECALHO}"
    planilha.oddFooter.left.text = "AcadFlow"
    planilha.oddFooter.right.text = "Página &P de &N"


def resposta_planilha_alunos(linhas, marco, filtros=()):
    """Monta o arquivo .xlsx da lista filtrada."""
    workbook = Workbook()
    workbook.properties.creator = "AcadFlow"
    workbook.properties.title = "Alunos por marco acadêmico"
    planilha = workbook.active
    planilha.title = "Alunos"

    linhas = list(linhas)
    colunas = _colunas_do_marco(marco)
    _escrever_topo(planilha, len(linhas), filtros, colunas)
    _escrever_cabecalho(planilha, colunas)
    for deslocamento, linha in enumerate(linhas, start=1):
        _escrever_aluno(
            planilha,
            LINHA_CABECALHO + deslocamento,
            _valores_do_aluno(linha, marco),
            colunas,
        )
    _configurar_exibicao(planilha, colunas)

    arquivo = BytesIO()
    workbook.save(arquivo)
    response = HttpResponse(
        arquivo.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    data_exportacao = timezone.localdate().isoformat()
    response["Content-Disposition"] = (
        f'attachment; filename="alunos_por_marco_{marco}_{data_exportacao}.xlsx"'
    )
    return response
